from flask import Flask, request, render_template_string, send_file, flash, redirect, url_for
from datetime import datetime
from docx import Document
import os
from flask_login import login_required, current_user

from ung_dung.xu_ly_ai.tao_giao_trinh import tao_giao_trinh
from ung_dung.co_so_du_lieu.ket_noi import khoi_tao_db, db, login_manager
from ung_dung.co_so_du_lieu.mo_hinh import NguoiDung
from ung_dung.co_so_du_lieu.thao_tac_du_lieu import luu_lich_su_giao_trinh, lay_lich_su_nguoi_dung
from ung_dung.giao_dien_nguoi_dung.xac_thuc import xac_thuc

app = Flask(__name__, static_folder='static', static_url_path='/static')

# Khởi tạo database và login
# Khởi tạo database và login
khoi_tao_db(app)

# Tự động tạo admin nếu chưa có
with app.app_context():
    from ung_dung.admin_setup import init_admin_account
    init_admin_account()

# Đăng ký blueprints
app.register_blueprint(xac_thuc)
from ung_dung.giao_dien_nguoi_dung.quan_tri import quan_tri
app.register_blueprint(quan_tri, url_prefix='/admin')


@login_manager.user_loader
def load_user(user_id):
    return NguoiDung.query.get(int(user_id))


from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from ung_dung.cot_loi.ket_noi_ai import goi_ai
import re

def xuat_docx(chu_de, html):
    doc = Document()
    
    # 1. Cấu hình Style chuẩn (Times New Roman, 13pt)
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(13)
    
    # Heading 1 Style
    h1_style = doc.styles['Heading 1']
    h1_font = h1_style.font
    h1_font.name = 'Times New Roman'
    h1_font.size = Pt(16)
    h1_font.bold = True
    h1_font.color.rgb = RGBColor(0, 0, 0) # Black
    
    # Heading 2 Style
    h2_style = doc.styles['Heading 2']
    h2_font = h2_style.font
    h2_font.name = 'Times New Roman'
    h2_font.size = Pt(14)
    h2_font.bold = True
    h2_font.color.rgb = RGBColor(0, 0, 0)

    # Title
    main_title = doc.add_heading(f"GIÁO TRÌNH: {chu_de.upper()}", level=0)
    main_title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    doc.add_paragraph(f"Ngày tạo: {datetime.now().strftime('%d/%m/%Y')}", style='Normal').alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    doc.add_paragraph("") # Spacing

    # Helper: Cleaning only (Translation already done in tao_giao_trinh)
    def clean_text(text):
        text = re.sub(r'<[^>]+>', '', text).strip()
        if not text: 
            return ""
        
        # Xử lý khoảng trắng và xuống dòng
        text = text.replace('\n', ' ').replace('\r', '')
        text = re.sub(r'\s+', ' ', text)
        return text

    # 2. Parsing HTML Structure
    patterns = [
        (r'<h3[^>]*class=["\']chuong-title["\'][^>]*>(.*?)</h3>', 'h1'), 
        (r'<h4[^>]*class=["\']sub-title["\'][^>]*>(.*?)</h4>', 'h2'),   
        (r'<p[^>]*class=["\']paragraph["\'][^>]*>(.*?)</p>', 'p'),       
        (r'<div[^>]*class=["\'](?:info-item|clo-item|practice-item|eval-item|ref-item)["\'][^>]*>(.*?)</div>', 'list')
    ]
    
    # Safe extraction helper
    def extract_content(m, tag_end):
         res = re.search(rf'>\s*(.*?)\s*</{tag_end}>', m, re.DOTALL)
         return res.group(1) if res else ""

    combined_pattern = r'(<h3[^>]*class=["\']chuong-title["\'][^>]*>.*?</h3>)|' \
                       r'(<h4[^>]*class=["\']sub-title["\'][^>]*>.*?</h4>)|' \
                       r'(<p[^>]*class=["\']paragraph["\'][^>]*>.*?</p>)|' \
                       r'(<div[^>]*class=["\'](?:info-item|clo-item|practice-item|eval-item|ref-item)["\'][^>]*>.*?</div>)'
                       
    matches = re.finditer(combined_pattern, html, re.DOTALL)
    
    for match in matches:
        full_match = match.group(0)

        if 'class="chuong-title"' in full_match or "class='chuong-title'" in full_match:
            content = extract_content(full_match, 'h3')
            # Title
            title = clean_text(content)
            if title:
                h = doc.add_heading(title, level=1)
                h.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
            
        elif 'class="sub-title"' in full_match or "class='sub-title'" in full_match:
            content = extract_content(full_match, 'h4')
            # Sub-title
            sub = clean_text(content)
            if sub:
                h = doc.add_heading(sub, level=2)
                h.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT

        elif 'class="paragraph"' in full_match or "class='paragraph'" in full_match:
            content = extract_content(full_match, 'p')
            if content:
                text = clean_text(content)
                if text:
                    p = doc.add_paragraph(text)
                    p.alignment = WD_PARAGRAPH_ALIGNMENT.JUSTIFY
                    p.paragraph_format.first_line_indent = Pt(13 * 2) 
                    p.paragraph_format.space_after = Pt(12)
            
        elif 'class="info-item"' in full_match or 'class="clo-item"' in full_match or 'class="practice-item"' in full_match:
            # List item
            content = re.sub(r'<[^>]+>', '', full_match).strip()
            if content:
                list_item = clean_text(content)
                doc.add_paragraph(list_item, style='List Bullet')

    os.makedirs("xuat_giao_trinh", exist_ok=True)
    filename = f"Giao_Trinh_{chu_de}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"
    filename = re.sub(r'[\\/*?:"<>|]', "", filename)
    path = f"xuat_giao_trinh/{filename}"
    doc.save(path)
    return path


@app.route("/", methods=["GET"])
def trang_chu():
    if current_user.is_authenticated and current_user.la_admin:
        return redirect(url_for('quan_tri.dashboard'))
        
    chu_de = request.args.get("chu_de")
    ket_qua = ""
    da_dang_nhap = current_user.is_authenticated
    ten_nguoi_dung = current_user.ho_ten if da_dang_nhap else ""

    if chu_de:
        try:
            ket_qua_obj = tao_giao_trinh(chu_de)
            ket_qua = ket_qua_obj["noi_dung"]
            chu_de_hien_thi = ket_qua_obj.get("chu_de_vi", chu_de)
            
            # Lưu lịch sử nếu đã đăng nhập
            if da_dang_nhap:
                luu_lich_su_giao_trinh(
                    current_user.id,
                    chu_de_hien_thi,
                    ket_qua
                )
        except Exception as e:
            ket_qua = f"<div class='error-message'><p><b>⚠ Lỗi: {str(e)}</b></p></div>"

    return render_template_string(trang_chu_template(ket_qua, da_dang_nhap, ten_nguoi_dung))


@app.route("/xuat-docx", methods=["GET", "POST"])
@login_required
def xuat():
    chu_de = request.args.get("chu_de") or request.form.get("chu_de")
    html_truc_tiep = request.form.get("html_content")

    if not chu_de:
        flash("Chưa có chủ đề", "error")
        return redirect(url_for("trang_chu"))

    try:
        if html_truc_tiep:
            # [WYSIWYG] Xuất trực tiếp từ nội dung người dùng thấy
            print(f"✅ Xuat Word tu noi dung HTML truc tiep cho: {chu_de}")
            duong_dan = xuat_docx(chu_de, html_truc_tiep)
        else:
            # Quy trình cũ: Tạo mới từ AI
            kq = tao_giao_trinh(chu_de)
            chu_de_vi = kq.get("chu_de_vi", chu_de)
            duong_dan = xuat_docx(chu_de_vi, kq["noi_dung"])

        # Cập nhật lịch sử với đường dẫn file
        if current_user.is_authenticated:
            lich_su = lay_lich_su_nguoi_dung(current_user.id, 1)
            if lich_su and (lich_su[0].chu_de == chu_de or (html_truc_tiep and len(lich_su) > 0)):
                lich_su[0].duong_dan_file = duong_dan
                lich_su[0].da_xuat_file = True
                db.session.commit()

        return send_file(
            duong_dan,
            as_attachment=True,
            download_name=f"{chu_de}.docx"
        )
    except Exception as e:
        import traceback
        traceback.print_exc()
        flash(f"Lỗi khi xuất file: {str(e)}", "error")
        return redirect(url_for("trang_chu"))


@app.route("/lich-su", methods=["GET"])
@login_required
def lich_su():
    """Trang lịch sử tạo giáo trình"""
    danh_sach = lay_lich_su_nguoi_dung(current_user.id, 50)
    return render_template_string(lich_su_template(danh_sach))


def trang_chu_template(ket_qua, da_dang_nhap, ten_nguoi_dung):
    """Template HTML cho trang chủ"""
    user_menu = f"""
        <div class="user-menu">
            <span>👤 {ten_nguoi_dung}</span>
            <a href="/lich-su" class="btn-history">📜 Lịch sử</a>
            <a href="/dang-xuat" class="btn-logout">Đăng xuất</a>
        </div>
    """ if da_dang_nhap else """
        <div class="user-menu">
            <a href="/dang-nhap" class="btn-login">Đăng nhập</a>
            <a href="/dang-ky" class="btn-register">Đăng ký</a>
        </div>
    """
    
    return f"""
<!DOCTYPE html>
<html lang="vi">
<head>
    <meta charset="utf-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Hệ thống soạn giáo trình AI</title>
    <link rel="stylesheet" href="/static/phong_cach/giao_dien.css">
    <link rel="stylesheet" href="https://cdnjs.cloudflare.com/ajax/libs/font-awesome/6.0.0/css/all.min.css">
</head>
<body>
    <header class="header">
        <div class="header-content">
            <h1><i class="fas fa-book-reader"></i> HỆ THỐNG SOẠN GIÁO TRÌNH AI</h1>
            <div class="header-actions" style="display: flex; align-items: center; gap: 15px;">
                <button onclick="toggleVietnamese()" class="btn-translate">
                    <i class="fas fa-language"></i> Dịch Tiếng Việt
                </button>
                {user_menu}
            </div>
        </div>
    </header>

    <!-- Hidden Google Translate Element -->
    <div id="google_translate_element" style="display:none"></div>

    <script type="text/javascript">
        function googleTranslateElementInit() {{
            new google.translate.TranslateElement({{
                pageLanguage: 'auto',
                includedLanguages: 'vi',
                layout: google.translate.TranslateElement.InlineLayout.SIMPLE,
                autoDisplay: false
            }}, 'google_translate_element');
        }}

        function toggleVietnamese() {{
            // Check if already translated
            if (document.cookie.indexOf('googtrans=/auto/vi') !== -1) {{
                // Clear cookie to revert
                document.cookie = 'googtrans=; expires=Thu, 01 Jan 1970 00:00:00 UTC; path=/;';
                location.reload();
            }} else {{
                // Set cookie to translate
                document.cookie = 'googtrans=/auto/vi; path=/;';
                location.reload();
            }}
        }}

        function xuatWordHienTai() {{
            const content = document.querySelector('.content').innerHTML;
            if (!content || content.trim().length < 100) {{
                alert('Chưa có nội dung để xuất!');
                return;
            }}
            
            // Tạo form ẩn để gửi POST
            const form = document.createElement('form');
            form.method = 'POST';
            form.action = '/xuat-docx';
            
            const htmlInput = document.createElement('input');
            htmlInput.type = 'hidden';
            htmlInput.name = 'html_content';
            htmlInput.value = content;
            
            const chuDeInput = document.createElement('input');
            chuDeInput.type = 'hidden';
            chuDeInput.name = 'chu_de';
            chuDeInput.value = new URLSearchParams(window.location.search).get('chu_de') || 'Giao_Trinh';
            
            form.appendChild(htmlInput);
            form.appendChild(chuDeInput);
            document.body.appendChild(form);
            form.submit();
        }}
    </script>
    <script type="text/javascript" src="//translate.google.com/translate_a/element.js?cb=googleTranslateElementInit"></script>

    <div class="container">
        <div class="form-container">
            <h2 style="margin-bottom: 20px; color: var(--primary-dark);">Bắt đầu soạn thảo</h2>
            <form method="get">
                <input type="text" name="chu_de" placeholder="Nhập chủ đề (ví dụ: Trí tuệ nhân tạo, Python, Marketing...)" value="{request.args.get('chu_de', '')}" required>
                <div style="margin-top: 20px;">
                    <button type="submit" class="btn-create"><i class="fas fa-rocket"></i> Tạo giáo trình</button>
                    {"<button type='button' onclick='xuatWordHienTai()' class='btn-export'><i class='fas fa-file-word'></i> Xuất Word</button>" if da_dang_nhap and request.args.get('chu_de') else ""}
                </div>
            </form>
            {"<p style='color: var(--text-muted); margin-top: 20px; font-size: 0.9em;'><i class='fas fa-info-circle'></i> Đăng nhập để lưu lịch sử và xuất file Word</p>" if not da_dang_nhap else ""}
        </div>

        <div class="content">
            {ket_qua}
        </div>
    </div>
</body>
</html>
"""


def lich_su_template(danh_sach):
    """Template HTML cho trang lịch sử"""
    items_html = ""
    if danh_sach:
        for item in danh_sach:
            items_html += f"""
            <div class="history-item" style="background: white; padding: 20px; border-radius: 12px; margin-bottom: 20px; box-shadow: 0 2px 5px rgba(0,0,0,0.05); border: 1px solid #e2e8f0;">
                <div class="history-header" style="display: flex; justify-content: space-between; margin-bottom: 10px;">
                    <h3 style="margin: 0; color: #4f46e5;">{item.chu_de}</h3>
                    <span class="date" style="color: #64748b; font-size: 0.9em;">{item.ngay_tao.strftime('%d/%m/%Y %H:%M')}</span>
                </div>
                <div class="history-info" style="margin-bottom: 15px; color: #334155;">
                    <span style="margin-right: 20px;"><i class="fas fa-chart-bar"></i> {item.do_dai_ky_tu:,} ký tự</span>
                    {"<span style='color: #10b981; font-weight: 600;'><i class='fas fa-check-circle'></i> Đã xuất file</span>" if item.da_xuat_file else "<span style='color: #f59e0b;'><i class='fas fa-clock'></i> Chưa xuất</span>"}
                </div>
                <div class="history-actions" style="display: flex; gap: 10px;">
                    <a href="/?chu_de={item.chu_de}" style="background: #e0e7ff; color: #4338ca; padding: 8px 16px; border-radius: 6px; text-decoration: none; font-weight: 500;">Xem lại</a>
                    {"<a href='#' style='background: #d1fae5; color: #065f46; padding: 8px 16px; border-radius: 6px; text-decoration: none; font-weight: 500;'>Tải file</a>" if item.da_xuat_file else ""}
                </div>
            </div>
            """
    else:
        items_html = "<div style='text-align: center; padding: 60px; color: #94a3b8;'><i class='fas fa-folder-open' style='font-size: 48px; margin-bottom: 20px; display: block;'></i><p>Chưa có lịch sử tạo giáo trình nào.</p></div>"
    
    return f"""
<!DOCTYPE html>
<html lang="vi">
<head>
    <meta charset="utf-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Lịch sử - Hệ thống giáo trình AI</title>
    <link rel="stylesheet" href="/static/phong_cach/giao_dien.css">
    <link rel="stylesheet" href="https://cdnjs.cloudflare.com/ajax/libs/font-awesome/6.0.0/css/all.min.css">
</head>
<body>
    <header class="header">
        <div class="header-content">
            <h1><i class="fas fa-history"></i> LỊCH SỬ TẠO GIÁO TRÌNH</h1>
            <a href="/" class="btn-back"><i class="fas fa-arrow-left"></i> Về trang chủ</a>
        </div>
    </header>
    <div class="container">
        <h2 style="margin-bottom: 30px; border-bottom: 2px solid #e2e8f0; padding-bottom: 10px;">Danh sách đã tạo ({len(danh_sach)})</h2>
        {items_html}
    </div>
</body>
</html>
"""


if __name__ == "__main__":
    app.run(debug=True)
