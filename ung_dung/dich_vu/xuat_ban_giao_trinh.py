from docx import Document
import os

def xuat_file_word(ten_bai, data_json):
    doc = Document()
    doc.add_heading(f'GIÁO TRÌNH: {ten_bai.upper()}', 0)
    
    mapping = {
        "Mục tiêu": data_json.get('muc_tieu', ''),
        "Kiến thức cốt lõi": data_json.get('kien_thuc_cot_loi', ''),
        "Phân tích chuyên sâu": data_json.get('phan_tich', ''),
        "Ví dụ minh họa": data_json.get('vi_du', ''),
        "Mở rộng": data_json.get('mo_rong', '')
    }
    
    for tieu_de, noi_dung in mapping.items():
        if noi_dung:  # Chỉ thêm nếu có nội dung
            doc.add_heading(tieu_de, level=1)
            doc.add_paragraph(str(noi_dung))
        
    # Tạo thư mục nếu chưa tồn tại
    thu_muc = "kho_du_lieu/file_xuat_ban"
    os.makedirs(thu_muc, exist_ok=True)
    
    path = os.path.join(thu_muc, f"{ten_bai}.docx")
    doc.save(path)
    return path