from docx import Document
from datetime import datetime
import os


def xuat_giao_trinh_docx(chu_de, noi_dung_html):
    document = Document()

    # Tiêu đề
    document.add_heading(f"GIÁO TRÌNH: {chu_de.upper()}", level=1)

    # Tách nội dung đơn giản (HTML → text)
    dong_noi_dung = noi_dung_html.replace("<br>", "\n") \
                                 .replace("</p>", "\n") \
                                 .replace("<p>", "") \
                                 .replace("<ul>", "") \
                                 .replace("</ul>", "") \
                                 .replace("<li>", "- ") \
                                 .replace("</li>", "\n") \
                                 .replace("<h2>", "\n") \
                                 .replace("</h2>", "\n") \
                                 .replace("<h3>", "\n") \
                                 .replace("</h3>", "\n") \
                                 .replace("<h4>", "\n") \
                                 .replace("</h4>", "\n") \
                                 .replace("<hr>", "\n")

    for dong in dong_noi_dung.split("\n"):
        document.add_paragraph(dong)

    # Thư mục xuất
    thu_muc = "xuat_giao_trinh"
    os.makedirs(thu_muc, exist_ok=True)

    ten_file = f"{chu_de}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"
    duong_dan = os.path.join(thu_muc, ten_file)

    document.save(duong_dan)
    return duong_dan
